from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUTPUT = r"E:\lang_web_app\documentations\qa-public-id-url-binding-test-guide.docx"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_margins(cell, top=100, start=120, bottom=100, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def keep_with_next(paragraph):
    paragraph.paragraph_format.keep_with_next = True


def add_bullet(doc, text, level=0):
    p = doc.add_paragraph(style="List Bullet" if level == 0 else "List Bullet 2")
    p.add_run(text)
    return p


def add_number(doc, text):
    p = doc.add_paragraph(style="List Number")
    p.add_run(text)
    return p


def add_code(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.3)
    p.paragraph_format.right_indent = Inches(0.3)
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(6)
    run = p.add_run(text)
    run.font.name = "Consolas"
    run._element.rPr.rFonts.set(qn("w:ascii"), "Consolas")
    run._element.rPr.rFonts.set(qn("w:hAnsi"), "Consolas")
    run.font.size = Pt(9)
    return p


def add_test_case(doc, tc_id, title, objective, preconditions, steps, expected, evidence):
    heading = doc.add_heading(f"{tc_id}  {title}", level=2)
    keep_with_next(heading)

    table = doc.add_table(rows=0, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.columns[0].width = Inches(1.35)
    table.columns[1].width = Inches(5.65)

    rows = [
        ("Objective", objective),
        ("Preconditions", preconditions),
        ("Steps", "\n".join(f"{i + 1}. {step}" for i, step in enumerate(steps))),
        ("Expected result", expected),
        ("Evidence", evidence),
        ("Result", "PASS / FAIL    Tester: __________    Date: __________"),
    ]

    for index, (label, value) in enumerate(rows):
        cells = table.add_row().cells
        cells[0].width = Inches(1.35)
        cells[1].width = Inches(5.65)
        cells[0].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        cells[1].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_shading(cells[0], "D9EAF7")
        for cell in cells:
            set_cell_margins(cell)
        label_run = cells[0].paragraphs[0].add_run(label)
        label_run.bold = True
        value_paragraph = cells[1].paragraphs[0]
        value_paragraph.paragraph_format.space_after = Pt(0)
        value_paragraph.add_run(value)

    doc.add_paragraph().paragraph_format.space_after = Pt(2)


doc = Document()
section = doc.sections[0]
section.top_margin = Inches(0.7)
section.bottom_margin = Inches(0.65)
section.left_margin = Inches(0.8)
section.right_margin = Inches(0.8)

styles = doc.styles
styles["Normal"].font.name = "Aptos"
styles["Normal"]._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
styles["Normal"]._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
styles["Normal"].font.size = Pt(10.5)
styles["Normal"].paragraph_format.space_after = Pt(6)
styles["Normal"].paragraph_format.line_spacing = 1.08

for style_name, size in (("Title", 24), ("Heading 1", 16), ("Heading 2", 12.5)):
    style = styles[style_name]
    style.font.name = "Aptos Display"
    style._element.rPr.rFonts.set(qn("w:ascii"), "Aptos Display")
    style._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos Display")
    style.font.size = Pt(size)
    style.font.color.rgb = RGBColor(0, 0, 0)
    style.font.bold = True

styles["Heading 1"].paragraph_format.space_before = Pt(12)
styles["Heading 1"].paragraph_format.space_after = Pt(6)
styles["Heading 1"].paragraph_format.keep_with_next = True
styles["Heading 2"].paragraph_format.space_before = Pt(9)
styles["Heading 2"].paragraph_format.space_after = Pt(4)
styles["Heading 2"].paragraph_format.keep_with_next = True

title = doc.add_paragraph(style="Title")
title.alignment = WD_ALIGN_PARAGRAPH.LEFT
title.add_run("Public ID URL Binding QA Test Guide")

subtitle = doc.add_paragraph()
subtitle.paragraph_format.space_after = Pt(14)
run = subtitle.add_run("Buildings and Containments")
run.bold = True
run.font.size = Pt(13)

intro = doc.add_paragraph()
intro.add_run("Purpose. ").bold = True
intro.add_run(
    "This guide helps a junior QA engineer understand and test the public ID URL binding feature. "
    "The feature hides internal Building BIN and Containment ID values from normal application URLs. "
    "It does not remove those identifiers from the database or from business information displayed on a page."
)

audience = doc.add_paragraph()
audience.add_run("Audience. ").bold = True
audience.add_run("QA engineers with about one year of testing experience who can use the application, browser address bar, developer tools, and basic SQL.")

doc.add_heading("What the feature does", level=1)
doc.add_paragraph(
    "A new UUID field named public_id identifies a Building or Containment in a browser URL. "
    "Laravel receives the UUID, finds the matching record, and then continues using the existing BIN or Containment ID internally."
)

summary = doc.add_table(rows=1, cols=3)
summary.alignment = WD_TABLE_ALIGNMENT.CENTER
summary.autofit = False
widths = [Inches(1.7), Inches(2.65), Inches(2.65)]
for i, text in enumerate(("Item", "Before", "After")):
    cell = summary.rows[0].cells[i]
    cell.width = widths[i]
    set_cell_shading(cell, "1F4E78")
    set_cell_margins(cell)
    run = cell.paragraphs[0].add_run(text)
    run.bold = True
    run.font.color.rgb = RGBColor(255, 255, 255)
set_repeat_table_header(summary.rows[0])
examples = [
    ("Building edit", "/buildings/B016740/edit", "/buildings/4ee95bc0-.../edit"),
    ("Containment edit", "/containments/C011752/edit", "/containments/4294ec61-.../edit"),
    ("Containment map", "field=id and val=C011752", "field=public_id and val=4294ec61-..."),
]
for row_index, values in enumerate(examples):
    cells = summary.add_row().cells
    if row_index % 2:
        for cell in cells:
            set_cell_shading(cell, "F4F8FB")
    for i, value in enumerate(values):
        cells[i].width = widths[i]
        cells[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_margins(cells[i])
        cells[i].paragraphs[0].add_run(value)

doc.add_heading("Important terms", level=1)
terms = [
    ("BIN", "Internal Building identifier, such as B016740. It remains in the database and may still be displayed as business information."),
    ("Containment ID", "Internal containment identifier, such as C011752. It remains in database relationships and reports."),
    ("public_id", "A UUID used in normal browser URLs, such as 4294ec61-76b0-4f38-92ca-4ada119dcbf6."),
    ("URL binding", "Laravel uses the public_id from the URL to load the correct Building or Containment record."),
    ("404", "The expected Not Found response when the UUID is invalid, malformed, or does not match a record."),
]
for term, meaning in terms:
    p = doc.add_paragraph()
    p.add_run(f"{term}: ").bold = True
    p.add_run(meaning)

doc.add_heading("What QA should verify", level=1)
for item in [
    "Existing records have populated, valid, unique UUIDs.",
    "New records receive UUIDs automatically.",
    "Building URLs contain public_id and do not contain BIN.",
    "Containment URLs contain public_id and do not contain the internal Containment ID.",
    "A valid UUID opens the correct record.",
    "An invalid UUID or an old internal identifier does not open the record.",
    "Permissions still apply. Knowing a UUID must not grant unauthorized access.",
]:
    add_bullet(doc, item)

doc.add_paragraph().add_run(
    "Scope note: BIN and Containment ID may still appear inside page content, exports, reports, or search fields when they are required business identifiers. "
    "This test focuses on identifiers used in application URLs."
).italic = True

doc.add_page_break()
doc.add_heading("Test preparation", level=1)
doc.add_paragraph("Prepare the following before executing the test cases:")
for item in [
    "A QA environment with the public_id migration applied.",
    "A user who can view and edit Buildings and Containments.",
    "At least two Building records and two Containment records.",
    "Access to the browser address bar and, when available, PostgreSQL read-only access.",
    "The application base URL, for example http://127.0.0.1:8001.",
]:
    add_bullet(doc, item)

doc.add_heading("How to recognize a UUID", level=1)
doc.add_paragraph("A UUID normally has 36 characters, including four hyphens, in this shape:")
add_code(doc, "xxxxxxxx-xxxx-xxxx-xxxx-xxxxxxxxxxxx")
doc.add_paragraph("Example:")
add_code(doc, "4294ec61-76b0-4f38-92ca-4ada119dcbf6")
doc.add_paragraph("Do not accept a short sequential value such as B016740, C011752, 25, or 126 as a public ID.")

doc.add_heading("Manual test cases", level=1)

add_test_case(
    doc,
    "TC 01",
    "Building public IDs are populated unique UUIDs",
    "Confirm existing Buildings have valid and non-duplicated public IDs.",
    "The migration has run and QA has read-only database access.",
    [
        "Run the Building verification SQL shown below.",
        "Compare total, populated, and unique UUID counts.",
        "Check several public_id values visually against the UUID format.",
    ],
    "Total is greater than zero. Total, populated, and unique UUID counts are equal. Sample values use UUID format.",
    "Attach the SQL result or a screenshot with sensitive connection details hidden.",
)
add_code(doc, "SELECT count(*) AS total, count(public_id) AS populated,\n       count(DISTINCT public_id) AS unique_ids\nFROM building_info.buildings;")

add_test_case(
    doc,
    "TC 02",
    "Containment public IDs are populated unique UUIDs",
    "Confirm existing Containments have valid and non-duplicated public IDs.",
    "The migration has run and QA has read-only database access.",
    [
        "Run the Containment verification SQL shown below.",
        "Compare total, populated, and unique UUID counts.",
        "Check several public_id values visually against the UUID format.",
    ],
    "Total is greater than zero. Total, populated, and unique UUID counts are equal. Sample values use UUID format.",
    "Attach the SQL result or a screenshot with sensitive connection details hidden.",
)
add_code(doc, "SELECT count(*) AS total, count(public_id) AS populated,\n       count(DISTINCT public_id) AS unique_ids\nFROM fsm.containments;")

add_test_case(
    doc,
    "TC 03",
    "Building URLs use public ID and not BIN",
    "Verify Building actions do not expose BIN in the address bar.",
    "A Building with a known BIN is available and the tester has the required permissions.",
    [
        "Open the Building list.",
        "Record the BIN displayed for one Building.",
        "Open Detail, Edit, History, Connected Containments, Map, and Nearest Road actions when available.",
        "Inspect the address bar after each action.",
    ],
    "Each URL contains a UUID and does not contain the recorded BIN. Each action still works.",
    "Capture screenshots of the list record and each resulting address bar.",
)

add_test_case(
    doc,
    "TC 04",
    "Containment URLs use public ID and not internal ID",
    "Verify Containment actions do not expose the internal Containment ID in the address bar.",
    "A Containment with a known ID is available and the tester has the required permissions.",
    [
        "Open the Containment list.",
        "Record one displayed Containment ID.",
        "Open Detail, Edit, History, Type Change History, Connected Buildings, Map, Nearest Road, and Emptying History when available.",
        "Inspect the address bar after each action.",
    ],
    "Each record-specific URL contains a UUID and does not contain the recorded internal ID. Each action still works.",
    "Capture screenshots of the selected record and resulting address bars.",
)

add_test_case(
    doc,
    "TC 05",
    "Valid Building public ID resolves the correct record",
    "Confirm a Building UUID opens the Building it belongs to.",
    "QA can view the Building list and Building details.",
    [
        "Choose a Building and record its BIN, house number, and UUID from its Detail URL.",
        "Copy the complete URL.",
        "Open the copied URL in a new browser tab while signed in.",
        "Compare the displayed Building details with the original record.",
    ],
    "The copied UUID URL opens the same Building. BIN, house number, and other identifying details match.",
    "Capture the source list row and destination detail page.",
)

add_test_case(
    doc,
    "TC 06",
    "Valid Containment public ID resolves the correct record",
    "Confirm a Containment UUID opens the Containment it belongs to.",
    "QA can view the Containment list and Containment details.",
    [
        "Choose a Containment and record its displayed ID, type, and UUID from its Detail URL.",
        "Copy the complete URL.",
        "Open the copied URL in a new browser tab while signed in.",
        "Compare the displayed Containment details with the original record.",
    ],
    "The copied UUID URL opens the same Containment. ID, type, and other details match.",
    "Capture the source list row and destination detail page.",
)

add_test_case(
    doc,
    "TC 07",
    "Internal and invalid identifiers do not bind",
    "Confirm old identifiers, malformed values, and unknown UUIDs cannot be used as public route identifiers.",
    "A valid Building URL and valid Containment URL are available.",
    [
        "Replace the Building UUID in the URL with its BIN and load the page.",
        "Replace the Containment UUID in the URL with its internal ID and load the page.",
        "Replace a UUID with invalid-public-id and load the page.",
        "Replace a UUID with a correctly formatted but unknown UUID such as 00000000-0000-4000-8000-000000000000.",
    ],
    "Each request returns the normal 404 Not Found response. It must not return 500, expose a database error, or open another record.",
    "Capture the tested URLs and 404 responses. Record any 500 response as a high-priority defect.",
)

add_test_case(
    doc,
    "TC 08",
    "New models receive unique public IDs",
    "Confirm new Buildings and Containments receive UUIDs automatically.",
    "Use a QA database where creating records is allowed. Prepare valid test data.",
    [
        "Create two new Buildings using different test data.",
        "Open each new Building and record its UUID URL.",
        "Create two new Containments using different test data.",
        "Open each new Containment and record its UUID URL.",
        "Compare all generated UUIDs.",
    ],
    "Every new record has a valid UUID. No two tested records share the same UUID. URLs do not expose BIN or Containment ID.",
    "Record created test IDs and UUID URLs, then clean up test data according to the QA environment procedure.",
)

doc.add_page_break()
doc.add_heading("Additional security and regression checks", level=1)
doc.add_paragraph("Run these checks after the eight core cases:")
for item in [
    "Log in as a user without View Building permission and try a known Building UUID. Access must remain denied.",
    "Log in as a user without View Containment permission and try a known Containment UUID. Access must remain denied.",
    "Confirm Edit, Update, Delete, History, Map, Nearest Road, Connected Records, and Emptying History still work for authorized users.",
    "Check browser developer tools Network requests. Record-specific application requests should not expose BIN or Containment ID when the endpoint was included in this rollout.",
    "Confirm browser Back, Refresh, copied links, and bookmarks work with UUID URLs.",
]:
    add_bullet(doc, item)

doc.add_heading("Automated test command", level=1)
doc.add_paragraph("Developers or QA engineers with terminal access can run the existing automated test:")
add_code(doc, "cd E:\\lang_web_app\nphp artisan test tests/Feature/PublicIdUrlBindingTest.php")
doc.add_paragraph("A successful run displays PASS and eight passing tests. A failure displays FAIL, the test name, expected result, actual result, filename, and line number.")

doc.add_heading("Manual testing and automated testing", level=1)
comparison = doc.add_table(rows=1, cols=3)
comparison.alignment = WD_TABLE_ALIGNMENT.CENTER
comparison.autofit = False
comp_widths = [Inches(1.55), Inches(2.75), Inches(2.7)]
for i, value in enumerate(("Type", "Best for", "Stored in")):
    cell = comparison.rows[0].cells[i]
    cell.width = comp_widths[i]
    set_cell_shading(cell, "1F4E78")
    set_cell_margins(cell)
    run = cell.paragraphs[0].add_run(value)
    run.bold = True
    run.font.color.rgb = RGBColor(255, 255, 255)
set_repeat_table_header(comparison.rows[0])
comparison_rows = [
    ("Manual QA cases", "User behavior, UI links, permissions, screenshots, regression", "QA test plan or test management tool"),
    ("Automated feature tests", "UUID format, uniqueness, route generation, binding, repeatable regression", "tests/Feature/PublicIdUrlBindingTest.php"),
]
for row_index, values in enumerate(comparison_rows):
    cells = comparison.add_row().cells
    if row_index % 2:
        for cell in cells:
            set_cell_shading(cell, "F4F8FB")
    for i, value in enumerate(values):
        cells[i].width = comp_widths[i]
        cells[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_margins(cells[i])
        cells[i].paragraphs[0].add_run(value)

doc.add_heading("Defect reporting guide", level=1)
doc.add_paragraph("When a case fails, report the defect with enough information for a developer to reproduce it:")
for item in [
    "Short title, for example Containment Map URL exposes internal ID.",
    "Environment and application version or branch.",
    "User role and permissions.",
    "Record used, including BIN or Containment ID as test evidence.",
    "Exact steps and complete URL.",
    "Expected result and actual result.",
    "Screenshot, screen recording, or relevant Network response.",
    "Severity guidance: a 500/database error is higher priority than a cosmetic URL issue.",
]:
    add_bullet(doc, item)

doc.add_heading("QA execution summary", level=1)
results = doc.add_table(rows=1, cols=4)
results.alignment = WD_TABLE_ALIGNMENT.CENTER
results.autofit = False
result_widths = [Inches(0.8), Inches(3.55), Inches(1.1), Inches(1.55)]
for i, value in enumerate(("ID", "Scenario", "Result", "Defect ID")):
    cell = results.rows[0].cells[i]
    cell.width = result_widths[i]
    set_cell_shading(cell, "1F4E78")
    set_cell_margins(cell)
    run = cell.paragraphs[0].add_run(value)
    run.bold = True
    run.font.color.rgb = RGBColor(255, 255, 255)
set_repeat_table_header(results.rows[0])
scenario_names = [
    "Building UUID population and uniqueness",
    "Containment UUID population and uniqueness",
    "Building URLs hide BIN",
    "Containment URLs hide internal ID",
    "Building UUID resolves correct record",
    "Containment UUID resolves correct record",
    "Invalid and internal values return 404",
    "New records receive unique UUIDs",
]
for i, scenario in enumerate(scenario_names, 1):
    cells = results.add_row().cells
    if i % 2 == 0:
        for cell in cells:
            set_cell_shading(cell, "F4F8FB")
    values = (f"TC {i:02d}", scenario, "", "")
    for j, value in enumerate(values):
        cells[j].width = result_widths[j]
        cells[j].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_margins(cells[j], top=130, bottom=130)
        cells[j].paragraphs[0].add_run(value)
        if j in (0, 2, 3):
            cells[j].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER

footer = section.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
footer_run = footer.add_run("Public ID URL Binding QA Test Guide")
footer_run.font.size = Pt(8)
footer_run.font.color.rgb = RGBColor(90, 90, 90)

doc.core_properties.title = "Public ID URL Binding QA Test Guide"
doc.core_properties.subject = "Manual and automated QA scenarios for Building and Containment public IDs"
doc.core_properties.author = "IMIS Revamp Team"
doc.save(OUTPUT)
print(OUTPUT)
